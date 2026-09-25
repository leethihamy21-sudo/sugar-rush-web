"""Build-time: xuất final_scene.xlsx -> final_scene.json.

Chạy 1 lần trong Dockerfile (stage python) để runtime trên Render chỉ đọc
JSON nhẹ (~ms), không cần openpyxl/pandas lúc cold start.

Dùng stdlib + openpyxl (đã có trong requirements). Không import scene_1
để script chạy độc lập, không kéo theo Flask.
"""
import json
import sys

SRC = sys.argv[1] if len(sys.argv) > 1 else "final_scene.xlsx"
DST = sys.argv[2] if len(sys.argv) > 2 else "final_scene.json"
SHEET = sys.argv[3] if len(sys.argv) > 3 else "Chatbot Sugar Rush"


def main():
  from openpyxl import load_workbook

  wb = load_workbook(SRC, read_only=True, data_only=True)
  ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
  header = [
      str(c.value).strip() if c.value is not None else ""
      for c in next(ws.iter_rows(min_row=1, max_row=1))
  ]
  rows = []
  for r in ws.iter_rows(min_row=2, values_only=True):
    row = {header[i]: (r[i] if i < len(r) else "") for i in range(len(header))}
    intent = str(row.get("Intents", "") or "").strip()
    if not intent or intent.lower() == "nan":
      continue
    for k in ("Entities", "Intents", "Training", "Responses"):
      v = row.get(k)
      row[k] = "" if v is None else str(v)
    rows.append({k: row.get(k, "") for k in ("Entities", "Intents", "Training", "Responses")})
  wb.close()
  with open(DST, "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False)
  print(f"[export] {len(rows)} intents -> {DST}")


if __name__ == "__main__":
  main()
